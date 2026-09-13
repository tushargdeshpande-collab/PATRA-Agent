"""Extract plain text from uploaded PDF, DOCX or TXT files.

Every upload path in PATRA—resume, job description and supporting
evidence—passes through this module. Malformed or unsupported files
therefore fail with a clear error instead of crashing the application.
"""

from __future__ import annotations

import io
from pathlib import Path

from pypdf import PdfReader
from pypdf.errors import PdfReadError

from patra.config import MAX_UPLOAD_BYTES

try:
    import docx  # python-docx
except ImportError:  # pragma: no cover
    docx = None


class UnsupportedFileError(ValueError):
    """Raised when an uploaded file cannot be parsed."""


def _extract_pdf(data: bytes) -> str:
    """Extract selectable text from a PDF file."""

    try:
        reader = PdfReader(io.BytesIO(data))

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise UnsupportedFileError(
                    "This PDF is password-protected and cannot be read."
                ) from exc

        pages: list[str] = []

        for page in reader.pages:
            text = page.extract_text() or ""

            try:
                layout_text = (
                    page.extract_text(extraction_mode="layout") or ""
                )

                if len(layout_text) >= len(text) * 0.75:
                    text = layout_text

            except TypeError:
                # Older pypdf versions may not support extraction_mode.
                pass

            pages.append(text)

    except UnsupportedFileError:
        raise

    except PdfReadError as exc:
        raise UnsupportedFileError(
            "This PDF file appears to be corrupted or unreadable."
        ) from exc

    except Exception as exc:
        raise UnsupportedFileError(
            f"Could not read this PDF file ({exc})."
        ) from exc

    text = "\n".join(pages).strip()

    if not text:
        raise UnsupportedFileError(
            "No selectable text was found in this PDF. "
            "Scanned or image-only PDFs are not supported offline. "
            "Please upload a text-based PDF, DOCX or TXT file instead."
        )

    return text


def _extract_docx(data: bytes) -> str:
    """Extract paragraphs and table content from a DOCX file."""

    if docx is None:
        raise UnsupportedFileError(
            "DOCX support is unavailable because python-docx "
            "is not installed."
        )

    try:
        document = docx.Document(io.BytesIO(data))

    except Exception as exc:
        raise UnsupportedFileError(
            f"Could not read this DOCX file ({exc})."
        ) from exc

    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

            if row_text:
                parts.append(row_text)

    text = "\n".join(parts).strip()

    if not text:
        raise UnsupportedFileError(
            "This DOCX file does not contain any readable text."
        )

    return text


def _extract_txt(data: bytes) -> str:
    """Extract text using common text encodings."""

    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = data.decode(encoding).strip()

            if text:
                return text

        except (UnicodeDecodeError, LookupError):
            continue

    raise UnsupportedFileError(
        "This TXT file is empty or uses an unsupported text encoding."
    )


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_txt,
}


def extract_text(
    data: bytes,
    filename: str,
    allowed_extensions: set[str] | None = None,
) -> str:
    """Return text extracted from an uploaded PDF, DOCX or TXT file.

    UnsupportedFileError is raised for malformed, empty, oversized or
    unsupported files so the UI can display a readable error message.
    """

    if not data:
        raise UnsupportedFileError("The uploaded file is empty.")

    if len(data) > MAX_UPLOAD_BYTES:
        size_mb = MAX_UPLOAD_BYTES // (1024 * 1024)

        raise UnsupportedFileError(
            f"'{filename}' is larger than the {size_mb} MB upload limit."
        )

    suffix = Path(filename).suffix.lower()
    allowed = allowed_extensions or set(_EXTRACTORS)

    if suffix not in allowed:
        allowed_types = ", ".join(sorted(allowed))

        raise UnsupportedFileError(
            f"'{filename}' has an unsupported file type. "
            f"Allowed: {allowed_types}"
        )

    extractor = _EXTRACTORS.get(suffix)

    if extractor is None:
        raise UnsupportedFileError(
            f"No extractor is registered for '{suffix}' files."
        )

    return extractor(data)